"""Tests for nebula.ingest: parsing, sheet understanding and chunking.

All fixtures are generated in-test. reportlab's built-in fonts can't draw Devanagari,
so Hindi is covered through docx, markdown and CSV.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta

import pytest

from nebula.ingest import (
    UnsupportedFormat,
    chunk_document,
    classify_sheet,
    infer_column_types,
    normalize_column_names,
    parse,
    sheet_card,
)
from nebula.ingest.sheets import column_name_mapping, detect_header, to_date, to_number
from nebula.ingest.types import Block, Document, Sheet

HINDI = "आंगनवाड़ी केंद्रों में बच्चों की उपस्थिति बढ़ी। माता-पिता ने हर हफ्ते गतिविधियाँ कीं।"


# --------------------------------------------------------------------------- fixtures

def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Annual Plan 2024", 0)
    d.add_paragraph("This plan covers our work with Anganwadi workers.")
    d.add_heading("Andhra Pradesh", 1)
    d.add_paragraph("We work in three districts.")
    d.add_paragraph("Guntur", style="List Bullet")
    d.add_paragraph("Krishna", style="List Bullet")
    t = d.add_table(rows=3, cols=3)
    rows = [["District", "Centres", "Children"], ["Guntur", "120", "3,400"], ["Krishna", "95", "2,800"]]
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_heading("Guntur rollout", 2)
    d.add_paragraph(HINDI)
    d.add_heading("Uttar Pradesh", 1)
    d.add_paragraph("Pilot starts in March.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Programme overview"
    s1.placeholders[1].text = "Reach: 2 states"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Team"
    s2.placeholders[1].text = "Ops and content"
    s2.notes_slide.notes_text_frame.text = "Mention the new hires."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "District results"
    rows = [["District", "Centres"], ["Guntur", "120"], ["Krishna", "95"]]
    table = s3.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    box = s3.shapes.add_textbox(Inches(1), Inches(5), Inches(6), Inches(1))
    box.text_frame.text = (
        "Attendance improved across both districts after the parent WhatsApp nudges were "
        "introduced, with the largest gains in centres that had a dedicated supervisor visit."
    )
    s3.notes_slide.notes_text_frame.text = "Numbers are provisional until the Q3 audit."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["District attendance FY 2024-25"])
    ws.append([])
    ws.append(["District", "Block", "Visit date", "Children enrolled", "Attendance %", "Active"])
    districts = ["Guntur", "Krishna", "Prakasam"]
    start = datetime(2024, 4, 1)
    for i in range(30):
        ws.append([districts[i % 3], f"Block {i % 5}", start + timedelta(days=12 * i), 10 + i * 7,
                   round(50 + i * 1.3, 1), i % 4 != 0])

    faq = wb.create_sheet("FAQ")
    faq.append(["Question", "Answer", "Owner"])
    faq.append(["How do parents join the programme?",
                "Parents join by sending a hello message to the programme WhatsApp number. "
                "The bot replies with the registration steps in Hindi or English.", "Ops"])
    faq.append(["What if a parent has no smartphone?",
                "The Anganwadi worker shares the activity during the home visit. "
                "Printed activity cards are also available in every centre.", "Content"])
    faq.append(["Who approves new activities?",
                "The content team drafts them, the state lead reviews them and the "
                "programme head signs off before release.", "Content"])

    dash = wb.create_sheet("Dashboard")
    dash["B2"] = "Programme dashboard"
    dash["B4"] = "Total enrolled"
    dash["D4"] = "=SUM(Attendance!D4:D33)"
    dash["B6"] = "Average attendance"
    dash["D6"] = "=AVERAGE(Attendance!E4:E33)"
    dash["F4"] = "Active centres"
    dash["H4"] = "=COUNTIF(Attendance!F4:F33,TRUE)"
    dash["B9"] = "Districts"
    dash["D9"] = "=COUNTA(Attendance!A4:A33)"

    hidden = wb.create_sheet("Lookups")
    hidden.append(["code", "name"])
    hidden.sheet_state = "hidden"

    merged = wb.create_sheet("Enrolment")
    merged.append(["District", "Enrolment", None])
    merged.append([None, "Boys", "Girls"])
    merged.merge_cells("A1:A2")
    merged.merge_cells("B1:C1")
    for name, b, g in [("Guntur", 120, 130), ("Krishna", 90, 101), ("Prakasam", 70, 75), ("Nellore", 60, 66)]:
        merged.append([name, b, g])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pdf(pages: list[list[str]]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    for lines in pages:
        y = 800
        for line in lines:
            if line:
                c.drawString(60, y, line)
            y -= 18
        c.showPage()
    c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- DOCX

def test_docx_order_headings_lists_tables_and_hindi():
    doc = parse(make_docx(), filename="Annual Plan.docx")
    assert doc.kind == "doc"
    assert doc.title == "Annual Plan"
    kinds = [(b.kind, b.level) for b in doc.blocks]
    assert kinds == [
        ("heading", 1), ("paragraph", None), ("heading", 1), ("paragraph", None), ("list", None),
        ("table", None), ("heading", 2), ("paragraph", None), ("heading", 1), ("paragraph", None),
    ]
    lst = doc.blocks[4]
    assert lst.text == "- Guntur\n- Krishna"
    table = doc.blocks[5]
    assert table.table == [["District", "Centres", "Children"], ["Guntur", "120", "3,400"], ["Krishna", "95", "2,800"]]
    assert table.text.startswith("| District | Centres | Children |\n| --- | --- | --- |")
    assert doc.blocks[7].text == HINDI
    assert doc.metadata["word_count"] > 20
    assert doc.metadata["table_count"] == 1


def test_docx_chunks_have_breadcrumbs_and_heading_paths():
    doc = parse(make_docx(), filename="Annual Plan.docx")
    chunks = chunk_document(doc)
    assert all(c.text.startswith("Annual Plan") for c in chunks)
    ap = next(c for c in chunks if "three districts" in c.text)
    assert ap.text.splitlines()[0] == "Annual Plan › Andhra Pradesh"
    assert ap.heading_path == ["Annual Plan", "Andhra Pradesh"]
    assert "| Guntur | 120 | 3,400 |" in ap.text  # the table stays with its section
    hindi = next(c for c in chunks if HINDI in c.text)
    assert hindi.heading_path == ["Annual Plan", "Andhra Pradesh", "Guntur rollout"]
    up = next(c for c in chunks if "Pilot starts" in c.text)
    assert up.heading_path == ["Annual Plan", "Uttar Pradesh"]


def test_docx_bold_lines_become_headings_when_no_heading_styles():
    import docx

    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Budget summary").bold = True
    d.add_paragraph("We spent less than planned.")
    buf = io.BytesIO()
    d.save(buf)
    doc = parse(buf.getvalue(), filename="notes.docx")
    assert doc.blocks[0].kind == "heading" and doc.blocks[0].level == 3


# --------------------------------------------------------------------------- PPTX

def test_pptx_slides_notes_and_tables():
    doc = parse(make_pptx(), filename="Q2 review.pptx")
    assert doc.kind == "slides" and doc.metadata["slide_count"] == 3
    assert [b.locator for b in doc.blocks] == [{"slide": 1}, {"slide": 2}, {"slide": 3}]
    s2 = doc.blocks[1].text
    assert s2.startswith("Slide 2: Team\n")
    assert "Ops and content" in s2 and s2.endswith("Notes: Mention the new hires.")
    s3 = doc.blocks[2].text
    assert "| District | Centres |" in s3 and "| Guntur | 120 |" in s3
    assert s3.index("| District") < s3.index("Attendance improved")  # top-to-bottom reading order
    assert "Notes: Numbers are provisional" in s3


def test_pptx_chunks_one_per_slide_tiny_slides_merged():
    chunks = chunk_document(parse(make_pptx(), filename="Q2 review.pptx"))
    assert all(c.kind == "slide" for c in chunks)
    assert chunks[0].locator == {"slides": [1, 2]}  # two tiny slides merged
    assert chunks[1].locator == {"slide": 3}
    assert chunks[1].heading_path == ["Q2 review", "District results"]
    assert chunks[1].text.startswith("Q2 review\nSlide 3: District results")


# --------------------------------------------------------------------------- XLSX

def test_xlsx_data_sheet_card_has_ranges_period_and_samples():
    doc = parse(make_xlsx(), filename="AP tracker.xlsx")
    assert doc.kind == "sheet"
    assert doc.metadata["hidden_sheets"] == ["Lookups"]
    assert doc.metadata["sheet_names"] == ["Attendance", "FAQ", "Dashboard", "Enrolment"]
    sheet = doc.sheets[0]
    assert sheet.kind == "data"
    assert sheet.header_row_index == 2
    assert sheet.header == ["District", "Block", "Visit date", "Children enrolled", "Attendance %", "Active"]
    assert sheet.n_rows == 30 and sheet.n_cols == 6
    assert "District attendance FY 2024-25" in sheet.notes
    assert infer_column_types(sheet) == ["text", "text", "date", "integer", "numeric", "boolean"]

    card = sheet_card(sheet, doc.title)
    assert len(card) <= 2500
    assert card.splitlines()[0] == "AP tracker › Attendance"
    assert "30 rows × 6 columns" in card
    assert "Period covered: 2024-04-01 to 2025-03-15 (Visit date)" in card
    assert "- Children enrolled (integer): 10 to 213" in card
    assert "Attendance % [attendance_pct] (numeric): 50 to 87.7" in card
    assert "District (text" in card and "Guntur" in card
    assert "Sample rows:" in card and "District: Guntur" in card

    chunks = [c for c in chunk_document(doc) if c.locator.get("sheet") == "Attendance"]
    assert len(chunks) == 1 and chunks[0].kind == "card"
    assert chunks[0].text == sheet_card(sheet, doc.title, max_chars=1800)


def test_xlsx_text_sheet_becomes_row_chunks():
    doc = parse(make_xlsx(), filename="AP tracker.xlsx")
    faq = next(s for s in doc.sheets if s.name == "FAQ")
    assert faq.kind == "text"
    rows = [c for c in chunk_document(doc) if c.locator.get("sheet") == "FAQ"]
    assert rows and all(c.kind == "row" for c in rows)
    assert rows[0].text.startswith("AP tracker › FAQ\nQuestion: How do parents join the programme?\nAnswer: ")
    assert rows[0].locator["rows"].startswith("2-")
    small = [c for c in chunk_document(doc, max_chars=300) if c.locator.get("sheet") == "FAQ"]
    assert [c.locator["rows"] for c in small] == ["2-2", "3-3", "4-4"]
    assert all(len(c.text) <= 300 for c in small)


def test_xlsx_dashboard_is_other_and_card_has_no_numbers():
    doc = parse(make_xlsx(), filename="AP tracker.xlsx")
    dash = next(s for s in doc.sheets if s.name == "Dashboard")
    assert dash.kind == "other" and dash.has_formulas
    card = next(c for c in chunk_document(doc) if c.locator.get("sheet") == "Dashboard")
    assert card.kind == "card"
    assert "Total enrolled" in card.text and "Average attendance" in card.text
    assert "Sample rows" not in card.text


def test_xlsx_two_row_merged_header_is_combined():
    doc = parse(make_xlsx(), filename="AP tracker.xlsx")
    s = next(s for s in doc.sheets if s.name == "Enrolment")
    assert s.header == ["District", "Enrolment - Boys", "Enrolment - Girls"]
    assert s.header_row_index == 1 and s.n_rows == 4 and s.kind == "data"
    assert normalize_column_names(s.header) == ["district", "enrolment_boys", "enrolment_girls"]


def test_classify_formula_heavy_and_layouts():
    header = ["Metric", "Q1", "Q2", "Q3"]
    rows = [["Enrolled", 10, 20, 30], ["Active", 5, 6, 7], ["Dropped", 1, 2, 3]]
    assert classify_sheet("Numbers", header, rows, has_formulas=False) == "data"
    assert classify_sheet("Numbers", header, rows, has_formulas=0.8) == "other"
    assert classify_sheet("Summary", header, rows, has_formulas=True) == "other"
    assert classify_sheet("Pivot", ["Row Labels", "Sum of Children"], [["AP", 3], ["UP", 4], ["Grand Total", 7]],
                          False) == "other"
    stacked = rows + [[None] * 4, ["District", "Jan", "Feb", "Mar"], ["Guntur", 1, 2, 3]]
    assert classify_sheet("Numbers", header, stacked, has_formulas=False) == "other"
    assert classify_sheet("Tiny", header, rows[:2], has_formulas=False) == "other"


def test_numeric_table_with_a_remarks_column_stays_data():
    header = ["District", "Centres", "Children", "Visits", "Remarks"]
    rows = [[f"D{i}", i, i * 10, i * 2, "Supervisor visited twice. Registers were up to date."] for i in range(10)]
    assert classify_sheet("Monthly", header, rows, has_formulas=False) == "data"


def test_wide_sheet_card_fits_budget_and_names_remaining_columns():
    header = [f"Indicator number {i} long name" for i in range(40)]
    sheet = Sheet("MIS", "data", header, [[j * i for i in range(40)] for j in range(50)], 50, 40)
    for limit in (2500, 1800, 600):
        card = sheet_card(sheet, "Monthly MIS", max_chars=limit)
        assert len(card) <= limit
        assert "Indicator number 39 long name" in card or "more: Indicator number" in card


# --------------------------------------------------------------------------- CSV

def test_csv_hindi_headers_and_indian_numbers():
    text = (
        "जिला,ब्लॉक,बच्चों की संख्या,Budget (₹),तारीख\n"
        "गुंटूर,तेनाली,\"1,23,456\",\"₹ 2,50,000\",01/04/2024\n"
        "कृष्णा,विजयवाड़ा,\"98,765\",\"₹ 1,75,500\",15/05/2024\n"
        "प्रकाशम,ओंगोल,\"4,321\",-,30/06/2024\n"
    )
    doc = parse(text.encode("utf-8-sig"), filename="जिला डेटा.csv")
    assert doc.title == "जिला डेटा"
    sheet = doc.sheets[0]
    assert sheet.header == ["जिला", "ब्लॉक", "बच्चों की संख्या", "Budget (₹)", "तारीख"]
    assert sheet.rows[0][0] == "गुंटूर"
    assert sheet.kind == "data"
    assert infer_column_types(sheet) == ["text", "text", "integer", "integer", "date"]
    names = normalize_column_names(sheet.header)
    assert names == ["col_1", "col_2", "col_3", "budget_inr", "col_5"]
    assert column_name_mapping(sheet.header)["col_1"] == "जिला"
    card = sheet_card(sheet, doc.title)
    assert "जिला [col_1] (text" in card and "गुंटूर" in card
    assert "बच्चों की संख्या [col_3] (integer): 4321 to 123456" in card
    assert "Period covered: 2024-04-01 to 2024-06-30 (तारीख)" in card


def test_csv_utf16_and_semicolons_and_tsv():
    doc = parse("name;score\nअ;1\nब;2\nक;3\n".encode("utf-16"), filename="scores.csv")
    assert doc.sheets[0].header == ["name", "score"] and doc.sheets[0].rows[1] == ["ब", "2"]
    doc = parse(b"a\tb\n1\t2\n3\t4\n5\t6\n", filename="x.tsv")
    assert doc.sheets[0].header == ["a", "b"] and doc.sheets[0].n_rows == 3
    doc = parse("name,city\nJos\xe9,M\xe1laga\n".encode("cp1252"), filename="x.csv")
    assert doc.sheets[0].rows[0] == ["José", "Málaga"]


# --------------------------------------------------------------------------- Markdown / HTML / text

MARKDOWN = f"""---
title: Field guide
---
# Field guide

Intro paragraph with **bold** text and an escaped \\- dash.

## Steps

- Visit the centre
- Meet the worker
  - Check the register
1. Numbered item

| District | Centres |
|---|---:|
| Guntur | 120 |
| Krishna | 95 |

> Parents are the first teachers.

```
code stays as is
```

![Map of districts][image1]

{HINDI}

[image1]: <data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAA>
"""


def test_markdown_blocks():
    doc = parse(MARKDOWN.encode(), mime="text/markdown", title="Guide")
    kinds = [b.kind for b in doc.blocks]
    assert kinds == ["heading", "paragraph", "heading", "list", "table", "quote", "code", "image", "paragraph"]
    assert doc.blocks[1].text == "Intro paragraph with bold text and an escaped - dash."
    assert doc.blocks[3].text == "- Visit the centre\n- Meet the worker\n  - Check the register\n1. Numbered item"
    assert doc.blocks[4].table == [["District", "Centres"], ["Guntur", "120"], ["Krishna", "95"]]
    assert doc.blocks[7].text == "Map of districts"
    assert doc.blocks[8].text == HINDI
    assert doc.metadata["front_matter"] == {"title": "Field guide"}
    assert not any("base64" in b.text for b in doc.blocks)
    assert doc.title == "Guide"
    assert parse(MARKDOWN.encode(), filename="guide.md").title == "guide"


def test_html_blocks_drop_scripts_and_nav():
    html = b"""<!doctype html><html><head><title>Centre report</title>
    <style>p{color:red}</style><script>var x = 'nope';</script></head><body>
    <nav><a href="/">Home</a> <a href="/about">About</a></nav>
    <h1>Centre report</h1><p>Visits went <b>well</b>.</p>
    <ul><li>Guntur<ul><li>Tenali</li></ul></li><li>Krishna</li></ul>
    <table><tr><th>District</th><th>Visits</th></tr><tr><td>Guntur</td><td>12</td></tr></table>
    <pre>line 1
line 2</pre><blockquote><p>Keep going.</p></blockquote><div>Loose div text</div>
    </body></html>"""
    doc = parse(html, filename="report.html")
    assert doc.kind == "html"
    assert [b.kind for b in doc.blocks] == ["heading", "paragraph", "list", "table", "code", "quote", "paragraph"]
    assert doc.blocks[1].text == "Visits went well."
    assert doc.blocks[2].text == "- Guntur\n  - Tenali\n- Krishna"
    assert doc.blocks[3].table == [["District", "Visits"], ["Guntur", "12"]]
    assert doc.blocks[4].text == "line 1\nline 2"
    joined = " ".join(b.text for b in doc.blocks)
    assert "nope" not in joined and "Home" not in joined and "color" not in joined
    assert parse(html).title == "Centre report"


def test_text_and_json():
    doc = parse(b"First para.\n\nSecond para.", filename="notes.txt")
    assert doc.kind == "text" and [b.text for b in doc.blocks] == ["First para.", "Second para."]
    doc = parse(b'{"state": "AP", "names": ["\xe0\xa4\x85"]}', filename="x.json")
    assert doc.blocks[0].kind == "code" and '"state": "AP"' in doc.blocks[0].text and "अ" in doc.blocks[0].text


# --------------------------------------------------------------------------- PDF

def test_pdf_pages_headings_and_locators():
    long_line = "Children attended sessions regularly across all centres in the district"
    pdf = make_pdf([
        ["1. Introduction", long_line, "and parents joined the weekly calls.", "",
         "Rocket Learning works with Anganwadi workers in many states", "to support early learning at home."],
        ["2. Findings", "Attendance rose by twelve percent compared with last year", "in every block we measured."],
    ])
    doc = parse(pdf, filename="report.pdf")
    assert doc.kind == "pdf" and doc.metadata["page_count"] == 2
    assert [(b.kind, b.locator) for b in doc.blocks] == [
        ("heading", {"page": 1}), ("paragraph", {"page": 1}), ("paragraph", {"page": 1}),
        ("heading", {"page": 2}), ("paragraph", {"page": 2}),
    ]
    assert doc.blocks[0].text == "1. Introduction" and doc.blocks[0].level == 1
    assert doc.blocks[1].text == f"{long_line} and parents joined the weekly calls."
    chunks = chunk_document(doc)
    assert [c.locator for c in chunks] == [{"page": 1}, {"page": 2}]
    assert chunks[1].text.startswith("report › 2. Findings\n")
    merged = chunk_document(Document("r", "pdf", [Block("paragraph", "a", locator={"page": 3}),
                                                  Block("paragraph", "b", locator={"page": 4})]))
    assert merged[0].locator == {"pages": [3, 4]}


def test_pdf_drops_running_headers_and_page_numbers_and_finds_bullets():
    bodies = ["Attendance grew in Guntur after the parent nudges started.",
              "Krishna saw the largest drop in absenteeism this quarter.",
              "Nellore centres reported better worker morale overall.",
              "Prakasam is the next district for the expansion."]
    pages = [
        ["Rocket Learning | Annual Report 2024", body, "- First point about centres",
         f"- Follow up in {body.split()[3]}", "", f"Page {i + 1} of 4"]
        for i, body in enumerate(bodies)
    ]
    doc = parse(make_pdf(pages), filename="annual.pdf")
    text = "\n".join(b.text for b in doc.blocks)
    assert "Annual Report" not in text and "Page 1 of 4" not in text
    assert all(body in text for body in bodies)
    lists = [b for b in doc.blocks if b.kind == "list"]
    assert len(lists) == 4 and lists[0].text == "- First point about centres\n- Follow up in Guntur"


def test_scanned_pdf_needs_ocr():
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.rect(50, 50, 400, 600, fill=1)  # an "image" with no text
    c.showPage()
    c.save()
    doc = parse(buf.getvalue(), filename="scan.pdf")
    assert doc.kind == "pdf" and doc.metadata["needs_ocr"] is True and not doc.blocks


# --------------------------------------------------------------------------- chunk sizes and splitting

def test_chunks_respect_max_chars_and_split_long_paragraphs_with_overlap():
    english = " ".join(f"Sentence number {i} talks about centre visits and parent engagement." for i in range(60))
    hindi = " ".join([HINDI] * 30)
    rows = [["District", "Block", "Centres", "Notes"]] + [
        [f"District {i}", f"Block {i}", str(i), "Supervisor visited and checked registers."] for i in range(80)
    ]
    doc = Document("Long report", "doc", [
        Block("heading", "Overview", level=1),
        Block("paragraph", english),
        Block("heading", "हिंदी सारांश", level=2),
        Block("paragraph", hindi),
        Block("heading", "Table", level=1),
        Block("table", "", table=rows),
        Block("list", "\n".join(f"- item {i} with some words to make it longer" for i in range(80))),
        Block("code", "\n".join(f"print({i})" for i in range(300))),
    ])
    chunks = chunk_document(doc, max_chars=600, overlap_chars=120)
    assert all(len(c.text) <= 600 for c in chunks)
    assert all(c.text.startswith("Long report") for c in chunks)
    eng = [c for c in chunks if "Sentence number" in c.text]
    assert len(eng) > 3
    # overlap: each continuation repeats the end of the previous piece
    first_body, second_body = eng[0].text.split("\n", 1)[1], eng[1].text.split("\n", 1)[1]
    assert second_body.split(" ", 1)[1][:30] in first_body or second_body[:30].lstrip("…") in first_body
    hi = [c for c in chunks if "आंगनवाड़ी" in c.text]
    assert len(hi) > 1 and all(c.heading_path == ["Long report", "Overview", "हिंदी सारांश"] for c in hi)
    assert all(c.text.rstrip().endswith("।") or c.text.rstrip().endswith("कीं।") for c in hi)
    tables = [c for c in chunks if "| District | Block | Centres | Notes |" in c.text]
    assert len(tables) > 1
    assert all("| --- |" in c.text for c in tables)  # header repeated in every piece
    for c in tables:  # no row is ever cut in half
        assert all(line.startswith("|") and line.endswith("|") for line in c.text.splitlines()[1:])
    code = [c for c in chunks if "print(" in c.text]
    assert all(c.text.split("\n", 1)[1].startswith("```") and c.text.endswith("```") for c in code)


def test_meeting_turns_are_windowed():
    turns = [Block("speaker_turn", f"Speaker {i % 2}: point number {i} about the rollout plan.",
                   locator={"t": f"00:{i:02d}:00"}) for i in range(40)]
    doc = Document("Weekly sync", "meeting", [Block("heading", "Transcript", level=1), *turns])
    chunks = chunk_document(doc, max_chars=400)
    assert all(c.kind == "turn" and len(c.text) <= 400 for c in chunks)
    assert chunks[0].locator["t_start"] == "00:00:00"
    assert chunks[-1].locator["t_end"] == "00:39:00"
    assert chunks[0].text.startswith("Weekly sync › Transcript\nSpeaker 0:")


# --------------------------------------------------------------------------- sheet helpers

def test_value_parsing():
    assert to_number("1,23,456") == 123456
    assert to_number("1,234,567") == 1234567
    assert to_number("12%") == 12.0
    assert to_number("₹ 2,50,000") == 250000
    assert to_number("(1,234)") == -1234
    assert to_number("12 34") is None and to_number("1,2") is None and to_number("abc") is None
    assert to_number(True) is None
    assert to_date("01/04/2024") == date(2024, 4, 1)
    assert to_date("April 2024") == date(2024, 4, 1)
    assert to_date("2024-04-01 10:30") == datetime(2024, 4, 1, 10, 30)
    assert to_date("Guntur district") is None and to_date("2024") is None


def test_detect_header_skips_title_rows():
    rows = [["Report for Q2", None, None], [None, None, None], ["District", "Centres", "Children"],
            ["Guntur", 12, 300], ["Krishna", 9, 250]]
    assert detect_header(rows) == 2
    assert detect_header([["District", 2022, 2023], ["Guntur", 10, 12], ["Krishna", 8, 9]]) == 0


def test_normalize_column_names():
    header = ["District Name", "district name", "2024 Target", "% Achieved", "Order", None, "x" * 100,
              "जिला / District", "समय"]
    names = normalize_column_names(header)
    assert names[:5] == ["district_name", "district_name_2", "c_2024_target", "pct_achieved", "order_"]
    assert names[5] == "col_6"
    assert len(names[6]) <= 60
    assert names[7] == "district" and names[8] == "col_9"
    assert len(set(names)) == len(names)


def test_infer_column_types_tolerates_blanks_and_markers():
    sheet = Sheet("s", "data", ["a", "b", "c", "d"], [
        [1, "1.5", None, "2024-01-01 09:00"],
        ["-", "2", "x", "2024-01-02 10:00"],
        ["3", "#N/A", "", None],
    ], 3, 4)
    assert infer_column_types(sheet) == ["integer", "numeric", "text", "timestamp"]


# --------------------------------------------------------------------------- robustness

def test_unsupported_and_unknown_inputs():
    with pytest.raises(UnsupportedFormat):
        parse(b"\x89PNG\r\n\x1a\n" + b"\x00" * 100, filename="photo.png")
    with pytest.raises(UnsupportedFormat):
        parse(b"whatever", filename="clip.mp4")
    assert parse(b"\x00\x01\x02\x03" * 100, filename="blob.bin").kind == "other"
    assert parse(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100, filename="old.xls").kind == "other"
    broken = parse(b"PK\x03\x04 not really a zip", filename="broken.docx")
    assert broken.kind == "other"
    assert parse(b"", filename="empty.docx").metadata == {"empty": True}
    # a zip that is a docx is detected by content even with a wrong name
    assert parse(make_docx(), filename="download").kind == "doc"


def test_document_round_trips_through_dict():
    doc = parse(make_docx(), filename="Annual Plan.docx")
    again = Document.from_dict(doc.to_dict())
    assert again == doc
