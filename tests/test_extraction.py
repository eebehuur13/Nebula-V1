"""What text extraction misses: legacy Hindi fonts, pictures and scans, and Claude filling the gaps."""

from __future__ import annotations

import base64
import io

from helpers import MEMBER, OUTSIDER, call, drain, raw_call
from nebula import access, db
from nebula.connectors.base import FetchedItem
from nebula.ingest import images, parse
from nebula.ingest.quality import assess
from nebula.knowledge import items


def _png(w: int = 800, h: int = 600) -> bytes:
    from PIL import Image

    img = Image.new("RGB", (w, h), (40, 120, 200))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def _docx(build) -> bytes:
    import docx

    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf_with_picture_page() -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas
    from PIL import Image

    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 750, "Centre visits in Guntur: the field team visited every centre in the district this quarter "
                          "and recorded attendance, nutrition and the state of each building in detail. " * 3)
    c.showPage()
    c.drawString(72, 750, "Attendance chart")
    c.drawImage(ImageReader(Image.open(io.BytesIO(_png()))), 72, 300, 400, 300)
    c.showPage()
    c.save()
    return buf.getvalue()


def _gaps(doc) -> dict[str, str]:
    return {g["part"]: g["state"] for g in assess(doc)["gaps"]}


def test_word_runs_in_kruti_dev_become_hindi():
    def build(d):
        p = d.add_paragraph()
        r = p.add_run("ftyk dk;ZØe vf/kdkjh")
        r.font.name = "Kruti Dev 010"
        d.add_paragraph("District programme officer (English stays as it is).")

    doc = parse(_docx(build), filename="letter.docx")
    texts = [b.text for b in doc.blocks]
    assert texts[0] == "जिला कार्यक्रम अधिकारी"
    assert texts[1] == "District programme officer (English stays as it is)."
    assert doc.metadata["legacy_fonts"]["kruti"] == ["Kruti Dev 010"]
    assert _gaps(doc) == {"all": "uncertain"}


def test_fonts_nebula_cannot_convert_are_flagged():
    def build(d):
        r = d.add_paragraph().add_run("Hkkjr ljdkj")
        r.font.name = "Walkman-Chanakya-905"

    doc = parse(_docx(build), filename="circular.docx")
    assert _gaps(doc) == {"all": "needs_extraction"}


def test_kruti_text_without_font_names_is_converted():
    text = ("# fjiksVZ\n\nftyk dk;ZØe vf/kdkjh }kjk lHkh vkaxuckM+h dsUnzksa dk fujh{k.k fd;k x;k vkSj fjiksVZ "
            "Hksth xbZA\n\nThe English summary stays in English.\n")
    doc = parse(text.encode(), filename="report.md")
    paragraphs = [b.text for b in doc.blocks if b.kind == "paragraph"]
    assert paragraphs[0].startswith("जिला कार्यक्रम अधिकारी द्वारा सभी")
    assert paragraphs[1] == "The English summary stays in English."


def test_google_doc_images_are_numbered_and_can_be_shown():
    png = base64.b64encode(_png(400, 300)).decode()
    md = (f"# Field visit\n\nWe visited the centre.\n\n![][image1]\n\nAttendance below.\n\n"
          f"[image1]: <data:image/png;base64,{png}>\n")
    doc = parse(md.encode(), filename="Field visit.md")
    assert [b.locator for b in doc.blocks if b.kind == "image"] == [{"image": 1}]
    assert _gaps(doc) == {"image:1": "uncertain"}
    pictures = images.render_part(md.encode(), "Field visit.md", "image:1")
    assert len(pictures) == 1 and pictures[0][1] == "png" and pictures[0][0][:4] == b"\x89PNG"


def test_word_images_small_logos_are_not_gaps():
    from docx.shared import Inches

    def build(d):
        d.add_picture(io.BytesIO(_png(100, 100)), width=Inches(0.6))   # a logo
        d.add_paragraph("Monthly report")
        d.add_picture(io.BytesIO(_png()), width=Inches(5))              # a chart
    data = _docx(build)
    doc = parse(data, filename="monthly.docx")
    assert _gaps(doc) == {"image:2": "uncertain"}
    assert images.render_part(data, "monthly.docx", "image:2")[0][1] == "png"


def test_pdf_picture_pages_and_slides_are_flagged():
    doc = parse(_pdf_with_picture_page(), filename="visits.pdf")
    assert _gaps(doc) == {"page:2": "needs_extraction"}
    page = images.render_part(_pdf_with_picture_page(), "visits.pdf", "page:2")
    assert page[0][1] == "png" and len(page[0][0]) > 1000

    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[5])
    s1.shapes.title.text = "Enrolment by district"
    s1.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(1.5), Inches(8), Inches(5))
    buf = io.BytesIO()
    prs.save(buf)
    deck = parse(buf.getvalue(), filename="review.pptx")
    assert _gaps(deck) == {"slide:1": "needs_extraction"}
    assert images.render_part(buf.getvalue(), "review.pptx", "slide:1")[0][1] == "png"


async def test_claude_reads_a_picture_page_and_saves_it_for_everyone(mcp):
    member = await access.ensure_person(MEMBER)
    nb = await db.fetch_one(
        """insert into collections (slug, name, kind, owner_person) values ('extraction-notes', 'Extraction notes',
           'notebook', %s) on conflict (slug) do update set name=excluded.name returning id""", (member.id,))
    fetched = FetchedItem(external_id="upload:visits-pdf", title="Guntur visits", url="", mime_type="application/pdf",
                          kind="upload", data=_pdf_with_picture_page(), filename="visits.pdf")
    async with db.tx() as conn:
        res = await items.upsert(conn, str(nb["id"]), fetched, created_by=member.id)
    await drain()
    item = f"item:{res['id']}"

    r = await call(mcp, MEMBER, "open", id=item)
    ext = r["extraction"]
    assert ext["status"] == "needs_extraction" and [g["part"] for g in ext["parts_not_read"]] == ["page:2"]
    assert "as_image" in ext["how_to_fix"]

    pic = await raw_call(mcp, MEMBER, "open", id=item, part="page:2", as_image=True)
    assert not pic.is_error and [c.type for c in pic.content] == ["text", "image"]
    assert pic.content[1].mime_type == "image/png"

    r = await call(mcp, MEMBER, "save", kind="transcription", id=item, part="page:2", expected_version=0,
                   body="Attendance chart: Guntur 92 percent, Krishna 88 percent.")
    assert "changed since you read it" in r["error"]
    r = await call(mcp, MEMBER, "save", kind="transcription", id=item, part="page:9", expected_version=1, body="x")
    assert "isn't a part of this file" in r["error"]
    r = await call(mcp, OUTSIDER, "save", kind="transcription", id=item, part="page:2", expected_version=1, body="x")
    assert "error" in r
    r = await call(mcp, MEMBER, "save", kind="transcription", id=item, part="page:2", expected_version=1,
                   body="Attendance chart: Guntur 92 percent, Krishna 88 percent.")
    assert r["saved"] and r["parts_still_not_read"] == []
    await drain()

    found = await call(mcp, MEMBER, "search", query="Krishna attendance 88 percent")
    assert any(x["id"] == item for x in found["results"])
    r = await call(mcp, MEMBER, "open", id=item)
    assert r["extraction"]["status"] == "transcribed" and "Read from pictures by Claude" in r["_body"]
    r = await call(mcp, MEMBER, "open", id=item, part="page:1")
    assert "Read from pictures" not in r["_body"]
    st = await call(mcp, MEMBER, "sources", action="status", id=item)
    assert "extraction" not in st or st["extraction"]["parts_not_read"] == 0
