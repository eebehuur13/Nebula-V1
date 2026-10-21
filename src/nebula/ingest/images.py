"""Pictures of the parts text extraction can't read, for Claude to look at.

`open` with `as_image=true` uses this: a PDF page is rendered as a picture; a slide gives the
pictures on it; `image:N` gives the Nth image in a Word file, a Google Doc (whose Markdown export
carries its images inline) or a workbook. Pictures are scaled down to at most 1600 pixels a side.
"""

from __future__ import annotations

import base64
import io
import re
import zipfile
from typing import Any

MAX_SIDE = 1600
MAX_IMAGES = 4
_DATA_URI = re.compile(r"data:(image/[\w.+-]+);base64,([A-Za-z0-9+/=\s]+)")
_MD_REF = re.compile(r"!\[[^\]]*\]\[([^\]]+)\]|!\[[^\]]*\]\((data:image/[^)]+)\)")
_MD_DEF = re.compile(r"^ {0,3}\[([^\]]+)\]:\s*<?(data:image/[^>\s]+)>?", re.M)


class NoImage(ValueError):
    pass


def _fit(data: bytes) -> tuple[bytes, str]:
    """Scale a picture down to MAX_SIDE and encode it as PNG (or JPEG for photos)."""
    from PIL import Image

    img = Image.open(io.BytesIO(data))
    img.load()
    if max(img.size) > MAX_SIDE:
        img.thumbnail((MAX_SIDE, MAX_SIDE))
    photo = img.mode in ("RGB", "CMYK", "YCbCr") and (img.format or "").upper() in ("JPEG", "JPG")
    if img.mode not in ("RGB", "RGBA", "L", "LA"):
        img = img.convert("RGBA" if "A" in img.getbands() else "RGB")
    out = io.BytesIO()
    if photo:
        img.convert("RGB").save(out, "JPEG", quality=85)
        return out.getvalue(), "jpeg"
    img.save(out, "PNG", optimize=True)
    return out.getvalue(), "png"


def _pdf_page(data: bytes, page: int) -> bytes:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        if not 1 <= page <= len(pdf):
            raise NoImage(f"The PDF has {len(pdf)} pages.")
        p = pdf[page - 1]
        width, height = p.get_size()
        scale = min(3.0, MAX_SIDE / max(width, height, 1))
        bitmap = p.render(scale=scale)
        out = io.BytesIO()
        bitmap.to_pil().save(out, "PNG")
        return out.getvalue()
    finally:
        pdf.close()


def _docx_images(data: bytes) -> list[bytes]:
    """Images in a Word file in reading order (the order the parser numbers them)."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8", "replace")
        targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        targets.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
        body = z.read("word/document.xml").decode("utf-8", "replace")
        out = []
        for drawing in re.findall(r"<wp:docPr\b.*?(?=<wp:docPr\b|$)", body, re.S):
            m = re.search(r'r:embed="([^"]+)"', drawing)
            target = targets.get(m.group(1), "") if m else ""
            path = "word/" + target.lstrip("/").removeprefix("word/") if target else ""
            out.append(z.read(path) if path in z.namelist() else b"")
        return out


def _xlsx_images(data: bytes) -> list[bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return [z.read(n) for n in sorted(z.namelist()) if n.startswith("xl/media/")]


def _markdown_images(text: str) -> list[bytes]:
    """Images of a Google Doc's Markdown export, in the order they appear in the text."""
    defs = {k: uri for k, uri in _MD_DEF.findall(text)}
    out = []
    for ref, inline in _MD_REF.findall(text):
        uri = inline or defs.get(ref, "")
        m = _DATA_URI.search(uri)
        out.append(base64.b64decode(re.sub(r"\s", "", m.group(2))) if m else b"")
    return out


def _slide_images(data: bytes, slide: int) -> list[bytes]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(io.BytesIO(data))
    if not 1 <= slide <= len(prs.slides):
        raise NoImage(f"The presentation has {len(prs.slides)} slides.")

    def pictures(shapes: Any) -> list[Any]:
        found = []
        for s in shapes:
            if getattr(s, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                found += pictures(s.shapes)
                continue
            try:
                found.append((s.width * s.height, s.image.blob))
            except Exception:
                pass
        return found

    return [blob for _, blob in sorted(pictures(prs.slides[slide - 1].shapes), key=lambda x: -(x[0] or 0))]


def render_part(data: bytes, filename: str, part: str) -> list[tuple[bytes, str]]:
    """Pictures for one part of a stored original: [(bytes, 'png' | 'jpeg'), ...]."""
    kind, _, value = part.partition(":")
    kind, value = kind.strip().lower(), value.strip()
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if not value.isdigit():
        raise NoImage("Give a page, slide or image number, e.g. 'page:3', 'slide:5' or 'image:2'.")
    n = int(value)
    if kind == "page" and (ext == "pdf" or data[:5] == b"%PDF-"):
        return [_fit(_pdf_page(data, n))]
    if kind == "slide" and ext == "pptx":
        blobs = _slide_images(data, n)
        if not blobs:
            raise NoImage(f"Slide {n} has no pictures; its text is already in the item.")
        return [_fit(b) for b in blobs[:MAX_IMAGES]]
    if kind == "image":
        if ext == "docx":
            images = _docx_images(data)
        elif ext in ("xlsx", "xlsm"):
            images = _xlsx_images(data)
        elif ext in ("md", "markdown"):
            images = _markdown_images(data.decode("utf-8", "replace"))
        else:
            raise NoImage("Numbered images can be shown for Word files, Google Docs and workbooks. "
                          "For a PDF, ask for the page ('page:3').")
        if not 1 <= n <= len(images) or not images[n - 1]:
            raise NoImage(f"There's no image {n} in this file ({len(images)} found).")
        return [_fit(images[n - 1])]
    raise NoImage(f"Nebula can't show '{part}' as a picture for this kind of file.")
