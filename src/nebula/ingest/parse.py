"""Turn file bytes into a structured `Document`.

`parse` sniffs the format (magic bytes first, then extension, then MIME type) and
hands off to one parser per format. Prose formats produce `Block`s in reading order;
spreadsheets produce `Sheet`s. Parsers never crash on messy input: a file that
can't be read comes back as kind "other" with `metadata["parse_error"]`. Only media
and archives raise `UnsupportedFormat`, so the caller can store them file-only.

Google exports arrive as Docs -> text/markdown (or docx), Sheets -> xlsx,
Slides -> pptx, and are handled by the same parsers.
"""

from __future__ import annotations

import codecs
import csv
import io
import json
import logging
import re
import warnings
import zipfile
from collections import Counter
from html.parser import HTMLParser
from collections.abc import Callable, Iterator
from typing import Any

from nebula.ingest.legacy_fonts import font_kind, kruti_to_unicode, looks_like_kruti
from nebula.ingest.sheets import classify_sheet, detect_header, format_value, is_blank
from nebula.ingest.types import Block, Document, Sheet

log = logging.getLogger(__name__)

__all__ = ["UnsupportedFormat", "markdown_table", "markdown_table_lines", "parse"]

MAX_SHEET_ROWS = 200_000
MAX_SHEET_COLS = 500


class UnsupportedFormat(ValueError):
    """Media, archives and other binaries we can't extract text from (store file-only)."""


# --------------------------------------------------------------------------- format detection

_EXT_FORMAT = {
    ".docx": "docx", ".docm": "docx", ".pptx": "pptx", ".pptm": "pptx", ".ppsx": "pptx",
    ".xlsx": "xlsx", ".xlsm": "xlsx", ".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".pdf": "pdf",
    ".md": "md", ".markdown": "md", ".txt": "txt", ".text": "txt", ".html": "html", ".htm": "html",
    ".json": "json",
}
_MIME_FORMAT = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "application/vnd.ms-excel.sheet.macroenabled.12": "xlsx",
    "text/csv": "csv", "application/csv": "csv", "text/tab-separated-values": "tsv",
    "application/pdf": "pdf", "text/markdown": "md", "text/x-markdown": "md", "text/plain": "txt",
    "text/html": "html", "application/xhtml+xml": "html", "application/json": "json",
    "application/vnd.google-apps.document": "md",
}
# Documents we recognise but can't read with our libraries: returned as kind "other".
_LEGACY_EXT = {".doc", ".xls", ".ppt", ".odt", ".ods", ".odp", ".rtf", ".pages", ".numbers", ".key"}
_MEDIA_EXT = {
    ".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff", ".heic", ".heif", ".svg", ".ico",
    ".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v", ".3gp", ".mp3", ".wav", ".m4a", ".aac", ".ogg",
    ".opus", ".flac", ".amr", ".zip", ".rar", ".7z", ".gz", ".tgz", ".tar", ".exe", ".apk", ".dmg", ".iso",
}
_MEDIA_MAGIC = (
    b"\x89PNG", b"\xff\xd8\xff", b"GIF87a", b"GIF89a", b"II*\x00", b"MM\x00*", b"ID3\x03", b"ID3\x04",
    b"OggS\x00", b"fLaC", b"\x1a\x45\xdf\xa3", b"Rar!\x1a", b"7z\xbc\xaf", b"\x1f\x8b", b"#!AMR\n",
)
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _ext(filename: str) -> str:
    name = filename.rsplit("/", 1)[-1]
    return "." + name.rsplit(".", 1)[-1].lower() if "." in name else ""


def _is_media(data: bytes) -> bool:
    head = data[:16]
    if head.startswith(_MEDIA_MAGIC):
        return True
    if head[:4] == b"RIFF" and head[8:12] in (b"WEBP", b"WAVE", b"AVI "):
        return True
    return head[:1] == b"\x00" and head[4:8] == b"ftyp"  # mp4 / mov / m4a / heic


def _zip_format(data: bytes) -> str | None:
    """docx/pptx/xlsx/legacy for a recognised zip, "archive" for any other zip, None if corrupt."""
    try:
        names = set(zipfile.ZipFile(io.BytesIO(data)).namelist())
    except zipfile.BadZipFile:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "ppt/presentation.xml" in names:
        return "pptx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    if "mimetype" in names:  # OpenDocument
        return "legacy"
    return "archive"


def _detect(data: bytes, filename: str, mime: str) -> str:
    ext = _ext(filename)
    if data.lstrip()[:5] == b"%PDF-" or b"%PDF-" in data[:1024]:
        return "pdf"
    is_zip = data[:4] == b"PK\x03\x04"
    if is_zip:
        fmt = _zip_format(data)
        if fmt == "archive":
            raise UnsupportedFormat(f"zip archive ({filename or mime or 'unnamed'})")
        if fmt:
            return fmt
    if _is_media(data) or ext in _MEDIA_EXT or mime.split("/")[0] in ("image", "video", "audio"):
        raise UnsupportedFormat(f"media file ({filename or mime or 'unnamed'})")
    if data[:8] == _OLE_MAGIC or ext in _LEGACY_EXT:
        return "legacy"
    fmt = _EXT_FORMAT.get(ext) or _MIME_FORMAT.get(mime)
    if fmt in ("docx", "pptx", "xlsx") and not is_zip:  # claimed Office but the bytes aren't a zip
        fmt = None
    if fmt:
        return fmt
    if _looks_like_text(data):
        head = _decode(data[:2000]).lstrip().lower()
        if head.startswith(("<!doctype html", "<html")):
            return "html"
        if head.startswith(("{", "[")):
            return "json"
        return "txt"
    return "other"


def _looks_like_text(data: bytes) -> bool:
    sample = data[:4096]
    if sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE, codecs.BOM_UTF8)):
        return True
    if b"\x00" in sample:
        return False
    text = _decode(sample)
    if not text:
        return False
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\r\t")
    return printable / len(text) > 0.95


def _decode(data: bytes, charset: str | None = None) -> str:
    """Bytes to text: BOMs, then the declared charset, UTF-8, BOM-less UTF-16, cp1252, latin-1."""
    if data.startswith(codecs.BOM_UTF8):
        return data[3:].decode("utf-8", errors="replace")
    if data.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE)):
        return data.decode("utf-16", errors="replace")
    if charset:
        try:
            return data.decode(charset)
        except (LookupError, UnicodeDecodeError):
            pass
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as e:
        if e.start >= len(data) - 3:  # a multi-byte char cut off at the end of a sample
            return data[: e.start].decode("utf-8")
    if data.count(b"\x00") > len(data) // 4:
        enc = "utf-16-le" if data[1::2].count(0) > data[0::2].count(0) else "utf-16-be"
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("cp1252")
    except UnicodeDecodeError:
        return data.decode("latin-1")


# --------------------------------------------------------------------------- entry point

def parse(data: bytes, filename: str = "", mime: str | None = None, title: str | None = None) -> Document:
    """Parse file bytes into a Document.

    The title is `title`, else the filename without a known extension, else whatever
    the file itself calls itself (docx Title, first heading, HTML <title>, PDF metadata).
    Raises `UnsupportedFormat` only for media/archives/executables.
    """
    mime_type, _, params = (mime or "").partition(";")
    mime_type = mime_type.strip().lower()
    m = re.search(r"charset=\"?([\w.-]+)", params, re.I)
    charset = m.group(1) if m else None
    base_title = (title or "").strip() or _title_from_filename(filename)

    if not data:
        return Document(title=base_title or "Untitled", kind="other", metadata={"empty": True})
    fmt = _detect(data, filename, mime_type)
    if fmt in ("legacy", "other"):
        meta: dict[str, Any] = {"format": fmt, "unsupported": True}
        return Document(title=base_title or "Untitled", kind="other", metadata=meta)

    try:
        doc = _PARSERS[fmt](data, charset)
    except UnsupportedFormat:
        raise
    except Exception as e:  # corrupt or exotic files must not take down a sync run
        log.warning("could not parse %r as %s: %s", filename, fmt, e)
        return Document(
            title=base_title or "Untitled",
            kind="other",
            metadata={"format": fmt, "parse_error": f"{type(e).__name__}: {e}"[:500]},
        )
    doc.title = base_title or doc.title.strip() or "Untitled"
    doc.metadata.setdefault("format", fmt)
    _convert_legacy_text(doc)
    if doc.blocks:
        doc.metadata.setdefault("word_count", sum(len(b.text.split()) for b in doc.blocks))
    return doc


_KEEP_PREFIX = re.compile(r"^(Slide \d+:? ?|Notes: |\s*[-*+•] |\s*\d{1,3}[.)] )")


def _convert_lines(text: str) -> str:
    """Kruti Dev -> Unicode, keeping Nebula's own prefixes ("Slide 3: ", "- ") intact."""
    out = []
    for line in text.split("\n"):
        m = _KEEP_PREFIX.match(line)
        keep = m.group(0) if m else ""
        out.append(keep + kruti_to_unicode(line[len(keep):]))
    return "\n".join(out)


def _convert_legacy_text(doc: Document) -> None:
    """Turn blocks that are really Hindi typed in Kruti Dev into Unicode (see legacy_fonts).

    Word and PDF files were already converted run by run where the font says Kruti Dev; this
    catches everything else (Google Docs exports, slides, HTML, text) by the text itself.
    """
    fonts = dict(doc.metadata.get("legacy_fonts") or {})
    font_seen = bool(fonts.get("kruti"))
    converted = 0
    for b in doc.blocks:
        if b.kind in ("image", "code"):
            continue
        if b.kind == "table" and b.table:
            cells = " ".join(c for row in b.table for c in row if c)
            if looks_like_kruti(cells, font_seen=font_seen):
                b.table = [[kruti_to_unicode(c) if c else c for c in row] for row in b.table]
                b.text = markdown_table(b.table)
                converted += 1
            continue
        plain = "\n".join(line[len(m.group(0)):] if (m := _KEEP_PREFIX.match(line)) else line
                          for line in b.text.split("\n"))
        if looks_like_kruti(plain, font_seen=font_seen):
            b.text = _convert_lines(b.text)
            converted += 1
    if converted:
        fonts["converted_blocks"] = fonts.get("converted_blocks", 0) + converted
    if fonts:
        doc.metadata["legacy_fonts"] = fonts


def _font_note(meta: dict[str, Any], kind: str | None, name: str | None) -> None:
    """Remember a legacy Hindi font seen in the file (kind 'kruti' is converted, 'legacy' isn't)."""
    if kind and name:
        fonts = meta.setdefault("legacy_fonts", {})
        bucket = fonts.setdefault("kruti" if kind == "kruti" else "unconverted", [])
        clean = str(name).lstrip("/").rsplit("+", 1)[-1]
        if clean not in bucket:
            bucket.append(clean)


def _title_from_filename(filename: str) -> str:
    name = filename.rsplit("/", 1)[-1].strip()
    ext = _ext(name)
    if ext and (ext in _EXT_FORMAT or ext in _LEGACY_EXT or ext in _MEDIA_EXT):
        name = name[: -len(ext)]
    return name.strip()


# --------------------------------------------------------------------------- shared helpers

def _clean(text: str) -> str:
    return " ".join(text.split())


def markdown_table_lines(rows: list[list[str]]) -> list[str]:
    """[header line, separator line, *row lines] of a pipe table; cells flattened to one line."""
    width = max((len(r) for r in rows), default=0)

    def line(r: list[str]) -> str:
        cells = [_clean(str(c)).replace("|", "\\|") for c in r] + [""] * (width - len(r))
        return "| " + " | ".join(cells) + " |"

    return [line(rows[0]), "| " + " | ".join(["---"] * width) + " |", *(line(r) for r in rows[1:])]


def markdown_table(rows: list[list[str]]) -> str:
    return "\n".join(markdown_table_lines(rows)) if rows else ""


def _table_blocks(rows: list[list[str]], locator: dict[str, Any] | None = None) -> list[Block]:
    """A table block, or plain paragraphs for layout tables with a single column."""
    rows = [[_clean(c) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return []
    loc = dict(locator or {})
    if all(sum(1 for c in r if c) <= 1 for r in rows):  # text boxes / one-column layouts
        return [Block("paragraph", " ".join(c for c in r if c), locator=dict(loc)) for r in rows]
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    return [Block("table", markdown_table(rows), locator=loc, table=rows)]


# --------------------------------------------------------------------------- DOCX

LARGE_IMAGE_EMU2 = 4 * 10**12   # about 2.5 x 1.75 inches: bigger than a logo or an icon

def _parse_docx(data: bytes, charset: str | None = None) -> Document:
    import docx
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run

    d = docx.Document(io.BytesIO(data))
    P, TBL, R = qn("w:p"), qn("w:tbl"), qn("w:r")
    SDT, SDT_CONTENT, CUSTOM_XML = qn("w:sdt"), qn("w:sdtContent"), qn("w:customXml")
    DOC_PR, EXTENT, RFONTS = qn("wp:docPr"), qn("wp:extent"), qn("w:rFonts")
    style_cache: dict[str | None, tuple[int | None, bool]] = {}
    font_cache: dict[str | None, str | None] = {}
    meta: dict[str, Any] = {}

    def rfont(rpr: Any) -> str | None:
        f = rpr.find(RFONTS) if rpr is not None else None
        return (f.get(qn("w:ascii")) or f.get(qn("w:hAnsi"))) if f is not None else None

    try:
        defaults = d.styles.element.find(qn("w:docDefaults"))
        rpr_default = defaults.find(qn("w:rPrDefault")) if defaults is not None else None
        default_font = rfont(rpr_default.find(qn("w:rPr"))) if rpr_default is not None else None
    except Exception:
        default_font = None

    def para_font(p: Any) -> str | None:
        sid = p.style
        if sid not in font_cache:
            name = None
            try:
                st = Paragraph(p, d._body).style
            except Exception:
                st = None
            for _ in range(5):
                if st is None:
                    break
                if st.font.name:
                    name = st.font.name
                    break
                st = st.base_style
            font_cache[sid] = name or default_font
        return font_cache[sid]

    def run_font(r: Any, p: Any) -> str | None:
        return rfont(r.find(qn("w:rPr"))) or para_font(p)

    def body_elements(parent: Any) -> Iterator[Any]:
        for child in parent.iterchildren():
            if child.tag in (P, TBL):
                yield child
            elif child.tag == SDT:
                content = child.find(SDT_CONTENT)
                if content is not None:
                    yield from body_elements(content)
            elif child.tag == CUSTOM_XML:
                yield from body_elements(child)

    def style_info(p: Any) -> tuple[int | None, bool]:
        """(heading level, is a list style) for a paragraph, following base styles."""
        sid = p.style
        if sid not in style_cache:
            level, is_list = None, False
            try:
                st = Paragraph(p, d._body).style
            except Exception:
                st = None
            for _ in range(5):
                if st is None:
                    break
                name = (st.name or "").strip().lower()
                if level is None:
                    m = re.fullmatch(r"heading\s*([1-6])", name)
                    level = 1 if name == "title" else (int(m.group(1)) if m else None)
                is_list = is_list or name.startswith("list")
                st = st.base_style
            style_cache[sid] = (level, is_list)
        return style_cache[sid]

    def par_text(p: Any) -> str:
        """The paragraph's text; runs typed in Kruti Dev are converted to Unicode Hindi."""
        parts = p.xpath("w:r | w:hyperlink | w:ins/w:r | w:smartTag/w:r | w:fldSimple/w:r | w:sdt/w:sdtContent/w:r")
        groups: list[list[Any]] = []  # [is kruti, text], consecutive runs merged so 'f' and 'Z' stay with their letters
        for e in parts:
            t = getattr(e, "text", "") or ""
            if not t:
                continue
            run = e if e.tag == R else next(e.iter(R), None)
            name = run_font(run, p) if run is not None else None
            kind = font_kind(name)
            _font_note(meta, kind, name)
            kruti = kind == "kruti" or (bool(groups) and groups[-1][0] and not t.strip())
            if groups and groups[-1][0] == kruti:
                groups[-1][1] += t
            else:
                groups.append([kruti, t])
        return "".join(kruti_to_unicode(t) if k else t for k, t in groups)

    def all_bold(p: Any) -> bool:
        runs = [r for r in p.iter(qn("w:r")) if (r.text or "").strip()]
        return bool(runs) and all(Run(r, None).bold for r in runs)  # type: ignore[arg-type]

    def table_rows(tbl: Any) -> list[list[str]]:
        rows: list[list[str]] = []
        above: dict[int, str] = {}  # last text per grid column, for vertically merged cells
        for tr in tbl.iterchildren(qn("w:tr")):
            row: list[str] = []
            col = 0
            for tc in tr.iterchildren(qn("w:tc")):
                span = max(1, tc.grid_span)
                if tc.vMerge == "continue":
                    text = above.get(col, "")
                else:
                    text = " ".join(t for t in (_clean(par_text(p)) for p in tc.iter(P)) if t)
                row.extend([text] + [""] * (span - 1))
                for k in range(span):
                    above[col + k] = text
                col += span
            rows.append(row)
        return rows

    blocks: list[Block] = []
    list_lines: list[str] = []
    bold_lines: list[int] = []
    images: list[int] = []
    large_images: list[int] = []

    def flush_list() -> None:
        if list_lines:
            blocks.append(Block("list", "\n".join(list_lines)))
            list_lines.clear()

    def add_images(drawings: list[Any]) -> None:
        for dp in drawings:
            images.append(len(images) + 1)
            extent = dp.getparent().find(EXTENT) if dp.getparent() is not None else None
            try:
                area = int(extent.get("cx", 0)) * int(extent.get("cy", 0)) if extent is not None else 0
            except ValueError:
                area = 0
            if area >= LARGE_IMAGE_EMU2:
                large_images.append(images[-1])
            alt = (dp.get("descr") or dp.get("title") or "").strip()
            blocks.append(Block("image", alt, locator={"image": images[-1]}))

    for el in body_elements(d.element.body):
        drawings = list(el.iter(DOC_PR))  # numbered in document order, as images.py finds them
        if el.tag == TBL:
            flush_list()
            blocks.extend(_table_blocks(table_rows(el)))
            add_images(drawings)
            continue
        text = par_text(el).strip()
        if text:
            level, list_style = style_info(el)
            numbered = el.xpath("./w:pPr/w:numPr")
            if level:
                flush_list()
                blocks.append(Block("heading", _clean(text), level=level))
            elif numbered or list_style:
                ilvl = el.xpath("./w:pPr/w:numPr/w:ilvl/@w:val")
                depth = int(ilvl[0]) if ilvl and str(ilvl[0]).isdigit() else 0
                list_lines.append("  " * min(depth, 6) + "- " + _clean(text))
            else:
                flush_list()
                if len(text) <= 100 and "\n" not in text and not text.endswith((".", ",", ";")) and all_bold(el):
                    bold_lines.append(len(blocks))
                blocks.append(Block("paragraph", text))
        if drawings:
            flush_list()
            add_images(drawings)
    flush_list()

    # Many real docs use bold lines instead of heading styles. Only when the document has no
    # styled headings at all, promote short all-bold lines to level-3 headings: they show up in
    # breadcrumbs but don't force chunk boundaries.
    if bold_lines and not any(b.kind == "heading" for b in blocks):
        for i in bold_lines:
            blocks[i] = Block("heading", _clean(blocks[i].text), level=3)

    meta.update({
        "heading_count": sum(b.kind == "heading" for b in blocks),
        "table_count": sum(b.kind == "table" for b in blocks),
    })
    if images:
        meta["images"] = {"count": len(images), "large": large_images, "sized": True}
    if bold_lines and meta["heading_count"] == len(bold_lines):
        meta["headings_from_bold"] = True
    cp = d.core_properties
    if cp.author:
        meta["author"] = cp.author
    if cp.modified:
        meta["modified"] = cp.modified.isoformat()
    first_heading = next((b.text for b in blocks if b.kind == "heading"), "")
    return Document(title=(cp.title or "").strip() or first_heading, kind="doc", blocks=blocks, metadata=meta)


# --------------------------------------------------------------------------- PPTX

def _parse_pptx(data: bytes, charset: str | None = None) -> Document:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    skip_placeholders = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER}
    prs = Presentation(io.BytesIO(data))

    def position(s: Any) -> tuple[int, int]:
        try:
            return (s.top or 0, s.left or 0)
        except Exception:
            return (0, 0)

    def shapes_in_order(shapes: Any) -> Iterator[Any]:
        for s in sorted(shapes, key=position):
            try:
                is_group = s.shape_type == MSO_SHAPE_TYPE.GROUP
            except Exception:
                is_group = False
            if is_group:
                yield from shapes_in_order(s.shapes)
            else:
                yield s

    def shape_lines(s: Any) -> list[str]:
        try:
            if s.is_placeholder and s.placeholder_format.type in skip_placeholders:
                return []
        except Exception:
            pass
        if getattr(s, "has_table", False):
            rows = [[_clean(c.text) for c in r.cells] for r in s.table.rows]
            return [markdown_table(rows)] if any(any(r) for r in rows) else []
        if getattr(s, "has_chart", False):
            return [_chart_line(s.chart)]
        if getattr(s, "has_text_frame", False):
            out = []
            for p in s.text_frame.paragraphs:
                t = p.text.replace("\v", "\n").strip()
                if t:
                    out.append("  " * (p.level or 0) + t)
            return out
        return []

    def is_picture(s: Any) -> bool:
        try:
            if s.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return True
        except Exception:
            pass
        try:
            s.image  # picture placeholders
            return True
        except Exception:
            return False

    def note_fonts(s: Any) -> None:
        if getattr(s, "has_text_frame", False):
            for p in s.text_frame.paragraphs:
                for r in p.runs:
                    _font_note(meta, font_kind(r.font.name), r.font.name)

    meta: dict[str, Any] = {}
    try:
        slide_area = int(prs.slide_width or 0) * int(prs.slide_height or 0)
    except Exception:
        slide_area = 0
    slide_area = slide_area or 12192000 * 6858000
    pictures: list[dict[str, int]] = []
    charts: list[int] = []
    blocks: list[Block] = []
    hidden: list[int] = []
    for n, slide in enumerate(prs.slides, 1):
        if slide._element.get("show") in ("0", "false"):
            hidden.append(n)
        title_shape = slide.shapes.title
        title = _clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        body: list[str] = []
        large = 0
        for s in shapes_in_order(slide.shapes):
            note_fonts(s)
            if is_picture(s):
                try:
                    if (s.width or 0) * (s.height or 0) >= 0.12 * slide_area:
                        large += 1
                except Exception:
                    pass
            if getattr(s, "has_chart", False):
                charts.append(n)
            if title_shape is not None and s.shape_id == title_shape.shape_id:
                continue
            body.extend(shape_lines(s))
        if large:
            pictures.append({"slide": n, "images": large, "words": len(" ".join([title, *body]).split())})
        if not title and body and len(body[0]) <= 80 and "\n" not in body[0] and not body[0].startswith("|"):
            title = body.pop(0).strip()
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        parts = [f"Slide {n}: {title}" if title else f"Slide {n}", *body]
        if notes:
            parts.append("Notes: " + notes)
        blocks.append(Block("slide", "\n".join(parts), locator={"slide": n}))

    meta["slide_count"] = len(blocks)
    if hidden:
        meta["hidden_slides"] = hidden
    if pictures:
        meta["slide_images"] = pictures
    if charts:
        meta["slide_charts"] = sorted(set(charts))
    title = (prs.core_properties.title or "").strip()
    return Document(title=title, kind="slides", blocks=blocks, metadata=meta)


def _chart_line(chart: Any) -> str:
    """One line describing a chart: its title and series names (values stay out of the index)."""
    try:
        title = _clean(chart.chart_title.text_frame.text) if chart.has_title else ""
    except Exception:
        title = ""
    try:
        series = [s.name for plot in chart.plots for s in plot.series if s.name][:8]
        cats = [str(c) for c in chart.plots[0].categories][:8] if chart.plots else []
    except Exception:
        series, cats = [], []
    line = "Chart" + (f": {title}" if title else "")
    extra = []
    if series:
        extra.append("series: " + ", ".join(series))
    if cats:
        extra.append("categories: " + ", ".join(cats))
    return line + (f" ({'; '.join(extra)})" if extra else "")


# --------------------------------------------------------------------------- PDF

_PAGE_NUMBER_RE = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*(?:of|/)\s*\d{1,4})?$|^[-–—]\s*\d{1,4}\s*[-–—]$", re.I)
_NUMBERED_HEADING_RE = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,3})\.?\s+(\S.*)$")
# Includes the private-use glyphs Word's Symbol/Wingdings bullets extract as.
_BULLET_RE = re.compile(r"^(?:[•●○◦▪▫■□➢➤►▶✓✔*\-–\uf0b7\uf0a7\uf076\uf0d8\uf0fc]|\(?\d{1,2}[.)]|\(?[a-z][.)])\s+")
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SMALL_WORDS = frozenset("a an and as at but by for from in into of on or the to vs via with".split())
_SENTENCE_END = (".", "?", "!", "।", "॥", ":", '"', "”")


def _pdf_heading_level(line: str) -> int | None:
    """Conservative heading guess for a PDF line: numbered ("2.1 Methods"), ALL CAPS or
    Title Case; short, at most 12 words, no trailing punctuation, no digits (unless numbered)."""
    s = line.strip()
    if not 3 <= len(s) <= 80 or s.endswith((".", ",", ";", ":", "।")) or len(s.split()) > 12:
        return None
    m = _NUMBERED_HEADING_RE.match(s)
    if m:
        rest = m.group(2)
        if rest[:1].isupper() and not any(ch.isdigit() for ch in rest) and len(rest.split()) <= 10:
            return min(m.group(1).count(".") + 1, 3)
        return None
    if any(ch.isdigit() for ch in s) or sum(ch.isalpha() for ch in s) < 3:
        return None
    words = s.split()
    if s.isupper() and len(words) <= 8 and (len(words) >= 2 or len(s) >= 4):
        return 1
    significant = [w for w in words if w[:1].isalpha() and w.lower() not in _SMALL_WORDS]
    if len(words) >= 2 and significant and all(w[0].isupper() for w in significant):
        return 2
    return None


def _edge_key(line: str) -> str:
    """Running header/footer key: digits normalised so "Page 3 of 9" matches every page.
    Long lines never count, so similar body sentences are safe."""
    return re.sub(r"\d+", "#", line.strip().lower()) if len(line) <= 70 else ""


def _join_lines(lines: list[str]) -> str:
    out = ""
    for line in lines:
        if out.endswith("-") and line[:1].islower():
            out = out[:-1] + line
        else:
            out = f"{out} {line}" if out else line
    return out


def _pdf_edges(lines: list[str]) -> tuple[set[str], set[str]]:
    """(page-number lines near the top/bottom, first and last content lines) of a page."""
    nonempty = [ln for ln in lines if ln]
    numbers = {ln for ln in nonempty[:2] + nonempty[-2:] if _PAGE_NUMBER_RE.match(ln)}
    content = [ln for ln in nonempty if ln not in numbers]
    return numbers, set(content[:1] + content[-1:])


def _pdf_page_blocks(lines: list[str], loc: dict[str, Any]) -> list[Block]:
    """Rebuild headings, bulleted lists and paragraphs from one page's text lines."""
    # A full line is about as long as the longer lines on the page; a line ending a sentence
    # well short of that is the last line of its paragraph.
    lengths = sorted(len(ln) for ln in lines if ln)
    typical = lengths[min(len(lengths) - 1, int(len(lengths) * 0.8))] if lengths else 0
    blocks: list[Block] = []
    para: list[str] = []
    items: list[str] = []
    item_open = False

    def flush() -> None:
        nonlocal item_open
        if para:
            blocks.append(Block("paragraph", _join_lines(para), locator=dict(loc)))
            para.clear()
        if items:
            blocks.append(Block("list", "\n".join(f"- {i}" for i in items), locator=dict(loc)))
            items.clear()
        item_open = False

    for line in lines:
        if not line:
            flush()
            continue
        at_boundary = (not para or para[-1].endswith(_SENTENCE_END)) and not item_open
        level = _pdf_heading_level(line) if at_boundary else None
        if level:
            flush()
            blocks.append(Block("heading", line, level=level, locator=dict(loc)))
            continue
        if _BULLET_RE.match(line):
            if para:
                flush()
            items.append(_BULLET_RE.sub("", line, count=1))
            item_open = True
        elif item_open:
            items[-1] = _join_lines([items[-1], line])
        else:
            if items:
                flush()
            para.append(line)
        if line.endswith(_SENTENCE_END) and len(line) < 0.75 * typical:
            if item_open:
                item_open = False
            else:
                flush()
    flush()
    return blocks


def _pdf_fonts(page: Any) -> list[tuple[str, str | None]]:
    """(font name, legacy kind) for each font the page uses."""
    out = []
    try:
        fonts = page["/Resources"].get_object().get("/Font")
        for ref in (fonts.get_object().values() if fonts else []):
            name = str(ref.get_object().get("/BaseFont", ""))
            out.append((name, font_kind(name)))
    except Exception:
        pass
    return out


def _pdf_page_text(page: Any, meta: dict[str, Any]) -> str:
    """The page's text, with anything set in Kruti Dev converted to Unicode Hindi."""
    fonts = _pdf_fonts(page)
    for name, kind in fonts:
        _font_note(meta, kind, name)
    if not any(kind == "kruti" for _, kind in fonts):
        return page.extract_text() or ""
    runs: list[tuple[str, bool]] = []

    def visit(text: str, cm: Any, tm: Any, font: Any, size: Any) -> None:
        name = str(font.get("/BaseFont", "")) if font else ""
        runs.append((text, font_kind(name) == "kruti"))

    text = page.extract_text(visitor_text=visit) or ""
    if all(kind == "kruti" for _, kind in fonts):
        return kruti_to_unicode(text)
    if "".join(t for t, _ in runs) == text:
        # Rebuild from the runs, converting only what was set in Kruti Dev.
        groups: list[list[Any]] = []
        for t, kruti in runs:
            kruti = kruti or (bool(groups) and groups[-1][0] and not t.strip())
            if groups and groups[-1][0] == kruti:
                groups[-1][1] += t
            else:
                groups.append([kruti, t])
        return "".join(kruti_to_unicode(t) if k else t for k, t in groups)
    for t in sorted({t.strip() for t, kruti in runs if kruti and len(t.strip()) >= 3}, key=len, reverse=True):
        text = text.replace(t, kruti_to_unicode(t))
    return text


def _pdf_large_images(node: Any, depth: int = 0) -> int:
    """How many images on a page (or a form inside it) are big enough to carry content, not a logo."""
    count = 0
    try:
        resources = node.get("/Resources")
        resources = resources.get_object() if resources is not None else None
        xobjects = resources.get("/XObject") if resources is not None else None
        for ref in (xobjects.get_object().values() if xobjects is not None else []):
            obj = ref.get_object()
            subtype = obj.get("/Subtype")
            if subtype == "/Image" and int(obj.get("/Width", 0)) * int(obj.get("/Height", 0)) >= 200_000:
                count += 1
            elif subtype == "/Form" and depth < 2:
                count += _pdf_large_images(obj, depth + 1)
    except Exception:
        pass
    return count


def _parse_pdf(data: bytes, charset: str | None = None) -> Document:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    meta: dict[str, Any] = {}
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            return Document(title="", kind="pdf", metadata={"encrypted": True})
    pages: list[str] = []
    page_images: list[dict[str, int]] = []
    for pno, page in enumerate(reader.pages, 1):
        try:
            pages.append(_pdf_page_text(page, meta))
        except Exception as e:  # one broken page shouldn't lose the rest
            log.debug("pdf page extract failed: %s", e)
            pages.append("")
        large = _pdf_large_images(page)
        if large:
            page_images.append({"page": pno, "images": large, "chars": len("".join(pages[-1].split()))})
    meta["page_count"] = len(pages)
    if page_images:
        meta["page_images"] = page_images
    try:
        info = reader.metadata
        doc_title = (info.title or "").strip() if info else ""
        if info and info.author:
            meta["author"] = str(info.author)
    except Exception:
        doc_title = ""

    empty_pages = [i + 1 for i, t in enumerate(pages) if len("".join(t.split())) < 20]
    if pages and len(empty_pages) >= 0.9 * len(pages):
        meta["needs_ocr"] = True
        return Document(title=doc_title, kind="pdf", metadata=meta)
    if empty_pages:
        meta["pages_without_text"] = empty_pages[:50]

    # A leading DEL is how bullets from fonts without a usable glyph map often come out.
    page_lines = [
        [_CONTROL_RE.sub("", re.sub(r"^\s*\x7f", "• ", ln)).strip() for ln in t.splitlines()] for t in pages
    ]
    # Running headers/footers: the first or last line on most pages (digits normalised).
    repeated: set[str] = set()
    if len(pages) >= 3:
        counts: Counter[str] = Counter()
        for lines in page_lines:
            counts.update({_edge_key(ln) for ln in _pdf_edges(lines)[1]})
        repeated = {k for k, c in counts.items() if k and c >= max(3, 0.6 * len(pages))}

    blocks: list[Block] = []
    for pno, lines in enumerate(page_lines, 1):
        numbers, outer = _pdf_edges(lines)
        kept = [ln for ln in lines if ln not in numbers and not (ln in outer and _edge_key(ln) in repeated)]
        blocks.extend(_pdf_page_blocks(kept, {"page": pno}))

    return Document(title=doc_title, kind="pdf", blocks=blocks, metadata=meta)


# --------------------------------------------------------------------------- Markdown

_MD_HEADING = re.compile(r"^ {0,3}(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^ {0,3}(`{3,}|~{3,})")
_MD_LIST = re.compile(r"^(\s*)([-*+•]|\d{1,3}[.)])\s+(.*)$")
_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_MD_REFDEF = re.compile(r"^ {0,3}\[[^\]]+\]:\s*\S")
_MD_HR = re.compile(r"^ {0,3}([-*_])(\s*\1){2,}\s*$")
_MD_SETEXT = re.compile(r"^ {0,3}(=+|-+)\s*$")
_MD_IMAGE = re.compile(r"!\[([^\]]*)\](?:\([^)]*\)|\[[^\]]*\])")
_MD_DATA_URI = re.compile(r"<?data:[\w/+.-]+;base64,[A-Za-z0-9+/=\s]*>?")
_MD_ESCAPE = re.compile(r"\\([\\`*_{}\[\]()#+\-.!|>~<])")
_MD_BR = re.compile(r"<br\s*/?>", re.I)


def _md_inline(text: str) -> str:
    """Tidy inline markdown for indexing: drop embedded images/data URIs and bold markers,
    undo backslash escapes (Google Docs exports escape a lot)."""
    text = _MD_DATA_URI.sub("", text)
    text = _MD_IMAGE.sub(lambda m: m.group(1), text)
    text = _MD_BR.sub(" ", text).replace("**", "").replace("__", "")
    return _MD_ESCAPE.sub(r"\1", text).strip()


def _md_cells(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [_md_inline(c) for c in re.split(r"(?<!\\)\|", s)]


def _markdown_blocks(text: str) -> tuple[list[Block], dict[str, Any]]:
    lines = text.splitlines()
    meta: dict[str, Any] = {}
    blocks: list[Block] = []
    i = 0
    if lines and lines[0].strip() == "---":  # YAML front matter
        end = next((j for j in range(1, min(len(lines), 200)) if lines[j].strip() in ("---", "...")), None)
        if end:
            try:
                import yaml

                fm = yaml.safe_load("\n".join(lines[1:end]))
                if isinstance(fm, dict):
                    meta["front_matter"] = {str(k): format_value(v, 300) for k, v in fm.items()}
            except Exception:
                pass
            i = end + 1

    para: list[str] = []
    pending_images: list[str] = []
    image_count = [0]

    def image(alt: str) -> Block:
        image_count[0] += 1
        return Block("image", alt.strip(), locator={"image": image_count[0]})

    def flush() -> None:
        if para:
            t = _md_inline(" ".join(para))
            if t:
                blocks.append(Block("paragraph", t))
            para.clear()
        while pending_images:
            blocks.append(image(pending_images.pop(0)))

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            flush()
            i += 1
            continue
        if m := _MD_FENCE.match(line):
            flush()
            fence, code = m.group(1), []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code.append(lines[i])
                i += 1
            i += 1
            if any(c.strip() for c in code):
                blocks.append(Block("code", "\n".join(code).strip("\n")))
            continue
        if m := _MD_HEADING.match(line):
            flush()
            t = _md_inline(m.group(2))
            if t:
                blocks.append(Block("heading", t, level=len(m.group(1))))
            i += 1
            continue
        if len(para) == 1 and _MD_SETEXT.match(line):
            t = _md_inline(para.pop())
            blocks.append(Block("heading", t, level=1 if stripped.startswith("=") else 2))
            i += 1
            continue
        if _MD_HR.match(line) or _MD_REFDEF.match(line):
            flush()
            i += 1
            continue
        if stripped.startswith("<!--"):
            flush()
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue
        if "|" in line and i + 1 < len(lines) and "-" in lines[i + 1] and _MD_TABLE_SEP.match(lines[i + 1]):
            flush()
            rows = [_md_cells(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(_md_cells(lines[i]))
                i += 1
            blocks.extend(_table_blocks(rows))
            continue
        if _MD_LIST.match(line):
            flush()
            items: list[str] = []
            while i < len(lines):
                cur = lines[i]
                if m := _MD_LIST.match(cur):
                    depth = len(m.group(1).expandtabs(4)) // 2
                    marker = m.group(2) if m.group(2)[0].isdigit() else "-"
                    items.append("  " * min(depth, 6) + f"{marker} {_md_inline(m.group(3))}")
                elif cur.strip() and items and cur[:1] in (" ", "\t"):
                    items[-1] += " " + _md_inline(cur.strip())
                elif not cur.strip():
                    j = i + 1
                    while j < len(lines) and not lines[j].strip():
                        j += 1
                    if j < len(lines) and _MD_LIST.match(lines[j]):
                        i = j
                        continue
                    break
                else:
                    break
                i += 1
            blocks.append(Block("list", "\n".join(items)))
            continue
        if stripped.startswith(">"):
            flush()
            quote: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            t = _md_inline("\n".join(quote))
            if t:
                blocks.append(Block("quote", t))
            continue
        if _MD_IMAGE.fullmatch(stripped):
            flush()
            blocks.append(image(_MD_IMAGE.fullmatch(stripped).group(1)))  # type: ignore[union-attr]
            i += 1
            continue
        pending_images.extend(m.group(1) for m in _MD_IMAGE.finditer(stripped))
        para.append(stripped)
        i += 1
    flush()
    if image_count[0]:
        meta["images"] = {"count": image_count[0], "large": []}
    return blocks, meta


def _parse_markdown(data: bytes, charset: str | None = None) -> Document:
    blocks, meta = _markdown_blocks(_decode(data, charset))
    fm_title = meta.get("front_matter", {}).get("title", "")
    first_h1 = next((b.text for b in blocks if b.kind == "heading" and b.level == 1), "")
    return Document(title=fm_title or first_h1, kind="doc", blocks=blocks, metadata=meta)


# --------------------------------------------------------------------------- HTML

_HTML_SKIP = {"script", "style", "nav", "noscript", "template", "svg", "iframe", "object"}
_HTML_BLOCK = {
    "p", "div", "section", "article", "main", "header", "footer", "aside", "figure", "figcaption",
    "dl", "dt", "dd", "form", "fieldset", "address", "details", "summary", "center", "hr", "body",
}


class _HTMLBlocks(HTMLParser):
    """Collects h1-h6, p, li, table, pre and blockquote into Blocks; drops script/style/nav."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[Block] = []
        self.title_parts: list[str] = []
        self.buf: list[str] = []
        self.skip = 0
        self.in_title = False
        self.heading: int | None = None
        self.list_depth = 0
        self.list_items: list[str] = []
        self.table_depth = 0
        self.rows: list[list[str]] = []
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.colspan = 1
        self.pre = 0
        self.quote_depth = 0
        self.quote_parts: list[str] = []

    def _flush(self) -> None:
        if self.pre:
            return
        text = _clean("".join(self.buf))
        self.buf.clear()
        if not text:
            return
        if self.cell is not None:
            self.cell.append(text)
        elif self.heading:
            self.blocks.append(Block("heading", text, level=self.heading))
        elif self.list_depth:
            self.list_items.append("  " * min(self.list_depth - 1, 6) + "- " + text)
        elif self.quote_depth:
            self.quote_parts.append(text)
        else:
            self.blocks.append(Block("paragraph", text))

    def _end_cell(self) -> None:
        if self.cell is not None and self.row is not None:
            self._flush()
            self.row.extend([" ".join(self.cell)] + [""] * (self.colspan - 1))
        self.cell = None

    def _end_row(self) -> None:
        self._end_cell()
        if self.row is not None:
            self.rows.append(self.row)
        self.row = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "title":
            self.in_title = True
            return
        if tag in _HTML_SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush()
            self.heading = int(tag[1])
        elif tag in ("ul", "ol"):
            self._flush()
            self.list_depth += 1
        elif tag == "li":
            self._flush()
        elif tag == "table":
            self._flush()
            self.table_depth += 1
            if self.table_depth == 1:
                self.rows = []
        elif tag == "tr" and self.table_depth == 1:
            self._end_row()
            self.row = []
        elif tag in ("td", "th") and self.table_depth == 1:
            self._end_cell()
            if self.row is None:
                self.row = []
            self.buf.clear()
            self.cell = []
            span = dict(attrs).get("colspan") or "1"
            self.colspan = int(span) if span.isdigit() and 0 < int(span) <= 50 else 1
        elif tag in ("td", "th", "tr"):  # nested table: keep cells apart
            self.buf.append(" ")
        elif tag == "pre":
            self._flush()
            self.pre += 1
        elif tag == "blockquote":
            self._flush()
            self.quote_depth += 1
        elif tag == "br":
            self.buf.append("\n" if self.pre else " ")
        elif tag == "img":
            alt = (dict(attrs).get("alt") or "").strip()
            if alt and self.cell is None:
                self._flush()
                self.blocks.append(Block("image", alt))
        elif tag in _HTML_BLOCK:
            self._flush()

    def handle_endtag(self, tag: str) -> None:
        if tag == "title":
            self.in_title = False
            return
        if tag in _HTML_SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag in ("h1", "h2", "h3", "h4", "h5", "h6"):
            self._flush()
            self.heading = None
        elif tag in ("ul", "ol"):
            self._flush()
            self.list_depth = max(0, self.list_depth - 1)
            if not self.list_depth and self.list_items:
                self.blocks.append(Block("list", "\n".join(self.list_items)))
                self.list_items = []
        elif tag == "li":
            self._flush()
        elif tag in ("td", "th") and self.table_depth == 1:
            self._end_cell()
        elif tag == "tr" and self.table_depth == 1:
            self._end_row()
        elif tag == "table" and self.table_depth:
            if self.table_depth == 1:
                self._end_row()
                self.blocks.extend(_table_blocks(self.rows))
                self.rows = []
            self.table_depth -= 1
        elif tag == "pre" and self.pre:
            self.pre -= 1
            if not self.pre:
                code = "".join(self.buf).strip("\n")
                self.buf.clear()
                if code.strip():
                    if self.cell is not None:
                        self.cell.append(_clean(code))
                    else:
                        self.blocks.append(Block("code", code))
        elif tag == "blockquote" and self.quote_depth:
            self._flush()
            self.quote_depth -= 1
            if not self.quote_depth and self.quote_parts:
                self.blocks.append(Block("quote", "\n".join(self.quote_parts)))
                self.quote_parts = []
        elif tag in _HTML_BLOCK:
            self._flush()

    def handle_data(self, data: str) -> None:
        if self.in_title:
            self.title_parts.append(data)
        elif not self.skip:
            self.buf.append(data)

    def close(self) -> None:
        super().close()
        self.pre = 0
        self._flush()
        if self.table_depth:
            self._end_row()
            self.blocks.extend(_table_blocks(self.rows))
        if self.list_items:
            self.blocks.append(Block("list", "\n".join(self.list_items)))
        if self.quote_parts:
            self.blocks.append(Block("quote", "\n".join(self.quote_parts)))


def _parse_html(data: bytes, charset: str | None = None) -> Document:
    if not charset:
        m = re.search(rb"<meta[^>]+charset=[\"']?([\w-]+)", data[:4096], re.I)
        charset = m.group(1).decode("ascii", "ignore") if m else None
    parser = _HTMLBlocks()
    parser.feed(_decode(data, charset))
    parser.close()
    return Document(title=_clean("".join(parser.title_parts)), kind="html", blocks=parser.blocks)


# --------------------------------------------------------------------------- plain text / JSON

def _text_blocks(text: str) -> list[Block]:
    paras = re.split(r"\n\s*\n", text.replace("\r\n", "\n").replace("\r", "\n"))
    return [Block("paragraph", p.strip()) for p in paras if p.strip()]


def _parse_text(data: bytes, charset: str | None = None) -> Document:
    return Document(title="", kind="text", blocks=_text_blocks(_decode(data, charset)))


def _parse_json(data: bytes, charset: str | None = None) -> Document:
    text = _decode(data, charset)
    try:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except ValueError:
        return Document(title="", kind="text", blocks=_text_blocks(text), metadata={"invalid_json": True})
    return Document(title="", kind="text", blocks=[Block("code", pretty)] if pretty.strip() else [])


# --------------------------------------------------------------------------- sheets

def _clean_cell(v: Any) -> Any:
    if isinstance(v, str):
        s = v.replace("\r\n", "\n").replace("\r", "\n").strip()
        return s or None
    if isinstance(v, float) and v != v:  # NaN
        return None
    return v


def _build_sheet(
    name: str,
    raw: list[list[Any]],
    *,
    formula_frac: float = 0.0,
    notes: list[str] | None = None,
    truncated: bool = False,
    header_spans: dict[int, list[tuple[int, int]]] | None = None,
    force_other: bool = False,
) -> Sheet:
    """Clean raw rows, find the header, classify, and wrap it all up as a Sheet.

    `header_spans` maps 0-based row index -> horizontally merged (first, last) 0-based
    column spans; used to combine a two-row header ("Enrolment" over "Boys", "Girls").
    """
    notes = list(notes or [])
    rows = [[_clean_cell(c) for c in r[:MAX_SHEET_COLS]] for r in raw]
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for j in range(len(r) - 1, width - 1, -1):
            if r[j] is not None:
                width = j + 1
                break
    if not width:
        return Sheet(name=name, kind="other", header=[], rows=[], n_rows=0, n_cols=0,
                     has_formulas=formula_frac > 0, notes=" ".join(notes + ["Empty sheet."]))
    rows = [(r + [None] * (width - len(r)))[:width] for r in rows]

    h = detect_header(rows)
    header_cells = list(rows[h])
    header_idx = h
    spans = (header_spans or {}).get(h)
    if spans and h + 1 < len(rows):
        sub = rows[h + 1]
        sub_cells = [c for c in sub if not is_blank(c)]
        if len(sub_cells) >= 2 and all(isinstance(c, str) for c in sub_cells):
            for c1, c2 in spans:
                for c in range(c1 + 1, min(c2 + 1, width)):
                    if is_blank(header_cells[c]):
                        header_cells[c] = header_cells[c1]
            combined = []
            for top, low in zip(header_cells, sub, strict=False):
                a = "" if is_blank(top) else format_value(top)
                b = "" if is_blank(low) else format_value(low)
                combined.append(a if not b or a == b else (b if not a else f"{a} - {b}"))
            header_cells = combined
            header_idx = h + 1
            notes.append(f"Two-row header (rows {h + 1}-{h + 2}) combined.")

    above = [format_value(c, 120) for r in rows[:h] for c in r if not is_blank(c)]
    if above:
        notes.insert(0, "Above the header: " + format_value("; ".join(above), 300))
    header = ["" if is_blank(c) else format_value(c) for c in header_cells]
    body = rows[header_idx + 1:]
    if len(body) > MAX_SHEET_ROWS:
        body, truncated = body[:MAX_SHEET_ROWS], True
    if truncated:
        notes.append(f"Truncated: only the first {MAX_SHEET_ROWS:,} rows were read.")
    kind = "other" if force_other else classify_sheet(name, header, body, formula_frac)
    return Sheet(
        name=name, kind=kind, header=header, rows=body, n_rows=len(body), n_cols=width,
        has_formulas=formula_frac > 0, header_row_index=header_idx, notes=" ".join(notes),
    )


_F_TAG = re.compile(rb"<(?:\w+:)?f[\s>/]")
_V_TAG = re.compile(rb"<(?:\w+:)?(?:v|is)[\s>]")
_MERGE_TAG = re.compile(rb'<(?:\w+:)?mergeCell\s+ref="([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)"')


def _col_index(letters: bytes) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + (ch - 64)
    return n


def _sheet_xml_facts(archive: zipfile.ZipFile | None, path: str | None) -> dict[str, Any]:
    """Formula share, merged ranges, pivots and drawings, read straight from the sheet XML
    (read-only openpyxl exposes none of these)."""
    facts: dict[str, Any] = {"formula_frac": 0.0, "merges": [], "pivot": False, "drawing": False}
    if archive is None or not path:
        return facts
    try:
        xml = archive.read(path.lstrip("/"))
    except KeyError:
        return facts
    n_f = len(_F_TAG.findall(xml))
    n_v = len(_V_TAG.findall(xml))
    facts["formula_frac"] = n_f / max(n_v, n_f, 1)
    facts["merges"] = [
        (_col_index(c1), int(r1), _col_index(c2), int(r2)) for c1, r1, c2, r2 in _MERGE_TAG.findall(xml)[:20000]
    ]
    folder, _, base = path.lstrip("/").rpartition("/")
    try:
        rels = archive.read(f"{folder}/_rels/{base}.rels")
        facts["pivot"] = b"/pivotTable" in rels
        facts["drawing"] = b"/drawing" in rels
    except KeyError:
        pass
    return facts


def _parse_xlsx(data: bytes, charset: str | None = None) -> Document:
    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets: list[Sheet] = []
    hidden: list[str] = []
    truncated_sheets: list[str] = []
    try:
        archive = getattr(wb, "_archive", None)
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                hidden.append(ws.title)
                continue
            facts = _sheet_xml_facts(archive, getattr(ws, "_worksheet_path", None))
            if hasattr(ws, "reset_dimensions"):
                ws.reset_dimensions()  # stored dimensions are often wrong; read what is there
            raw: list[list[Any]] = []
            truncated = False
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                for row in ws.iter_rows(values_only=True):
                    if len(raw) >= MAX_SHEET_ROWS + 20:
                        truncated = True
                        break
                    raw.append(list(row[:MAX_SHEET_COLS]))
            notes: list[str] = []
            header_spans: dict[int, list[tuple[int, int]]] = {}
            vertical = 0
            for c1, r1, c2, r2 in facts["merges"]:
                if r1 > len(raw):
                    continue
                if c2 > c1:
                    header_spans.setdefault(r1 - 1, []).append((c1 - 1, c2 - 1))
                if r2 > r1 and c1 - 1 < MAX_SHEET_COLS:
                    top = raw[r1 - 1][c1 - 1] if c1 - 1 < len(raw[r1 - 1]) else None
                    for r in range(r1, min(r2, len(raw))):
                        row = raw[r]
                        if len(row) < c1:
                            row.extend([None] * (c1 - len(row)))
                        row[c1 - 1] = top
                    vertical += 1
            if vertical:
                notes.append(f"{vertical} vertically merged cell range(s) filled down.")
            if facts["pivot"]:
                notes.append("Contains a pivot table.")
            if facts["drawing"]:
                notes.append("Has charts or images.")
            sheet = _build_sheet(
                ws.title, raw, formula_frac=facts["formula_frac"], notes=notes, truncated=truncated,
                header_spans=header_spans, force_other=facts["pivot"],
            )
            if "Truncated" in sheet.notes:
                truncated_sheets.append(ws.title)
            sheets.append(sheet)
        chart_sheets = [cs.title for cs in getattr(wb, "chartsheets", [])]
    finally:
        wb.close()
    meta: dict[str, Any] = {
        "sheet_names": [s.name for s in sheets],
        "sheet_count": len(sheets),
        "sheet_kinds": {s.name: s.kind for s in sheets},
        "row_counts": {s.name: s.n_rows for s in sheets},
    }
    if hidden:
        meta["hidden_sheets"] = hidden
    if chart_sheets:
        meta["chart_sheets"] = chart_sheets
    if truncated_sheets:
        meta["truncated_sheets"] = truncated_sheets
    return Document(title="", kind="sheet", sheets=sheets, metadata=meta)


def _sniff_delimiter(sample: str) -> str:
    """The delimiter that splits the first lines into the most consistent number (>1) of fields."""
    lines = sample.splitlines()[:50]
    best, best_score = ",", 0.0
    for delim in (",", "\t", ";", "|"):
        try:
            widths = [len(r) for r in csv.reader(lines, delimiter=delim) if r]
        except csv.Error:
            continue
        if not widths:
            continue
        mode, count = Counter(widths).most_common(1)[0]
        if mode < 2:
            continue
        score = count / len(widths) * min(mode, 50)
        if score > best_score:
            best, best_score = delim, score
    return best


def _parse_delimited(data: bytes, charset: str | None, delimiter: str | None) -> Document:
    text = _decode(data, charset)
    delimiter = delimiter or _sniff_delimiter(text[:50_000])
    csv.field_size_limit(max(csv.field_size_limit(), 10_000_000))
    raw: list[list[Any]] = []
    truncated = False
    for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
        if len(raw) >= MAX_SHEET_ROWS + 20:
            truncated = True
            break
        raw.append(row[:MAX_SHEET_COLS])
    sheet = _build_sheet("Sheet1", raw, truncated=truncated)
    meta: dict[str, Any] = {
        "sheet_names": [sheet.name],
        "sheet_count": 1,
        "sheet_kinds": {sheet.name: sheet.kind},
        "row_counts": {sheet.name: sheet.n_rows},
        "delimiter": delimiter,
    }
    if "Truncated" in sheet.notes:
        meta["truncated_sheets"] = [sheet.name]
    return Document(title="", kind="sheet", sheets=[sheet], metadata=meta)


_PARSERS: dict[str, Callable[[bytes, str | None], Document]] = {
    "docx": _parse_docx,
    "pptx": _parse_pptx,
    "xlsx": _parse_xlsx,
    "csv": lambda data, charset: _parse_delimited(data, charset, None),
    "tsv": lambda data, charset: _parse_delimited(data, charset, "\t"),
    "pdf": _parse_pdf,
    "md": _parse_markdown,
    "txt": _parse_text,
    "html": _parse_html,
    "json": _parse_json,
}
